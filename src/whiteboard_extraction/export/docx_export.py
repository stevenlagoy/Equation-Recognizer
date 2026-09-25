"""Export a BoardDocument to a Word (.docx) file."""

from pathlib import Path

from docx import Document

from whiteboard_extraction.reconstruction.document import BoardDocument


def export_docx(document: BoardDocument, output_path: Path) -> Path:
    """Write a BoardDocument to a .docx file with real, editable text."""
    doc = Document()
    doc.add_heading("Lecture Board Content", level=1)

    for block in document.blocks:
        minutes, seconds = divmod(int(block.timestamp), 60)
        heading = doc.add_paragraph()
        heading.add_run(f"{minutes:02d}:{seconds:02d}").bold = True
        doc.add_paragraph(block.text)

    doc.save(str(output_path))
    return output_path